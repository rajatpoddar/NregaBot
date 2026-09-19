# tests/test_skilled_aadhaar_report_tab.py
import pytest
from pathlib import Path


from src.tabs.skilled_aadhaar_report_tab import parse_skilled_aadhaar_table


def test_parse_skilled_aadhaar_table():
    fixture_path = Path("docs/htm/SkilledBulkAadhaarUpdate.aspx.html")
    assert fixture_path.exists(), "Snapshot fixture missing"
    
    html_content = fixture_path.read_text(encoding="utf-8")
    data, summary = parse_skilled_aadhaar_table(html_content)
    
    assert summary["total"] > 0
    assert summary["total"] == summary["male"] + summary["female"] + summary["other"]
    assert summary["total"] == summary["seeded"] + summary["pending"]
    assert len(data) == summary["total"]
    
    # Check specific fields of first row
    first = data[0]
    assert first["sno"] == "1"
    assert first["job_card_no"] == "S-JH-22-003-008-001/222288"
    assert first["worker_name"] == "MANOJ KUMAR SINGH"
    assert first["gender"] == "M"
    assert first["status"] in ("Seeded", "Pending")


def test_export_professional_report_excel(tmp_path, monkeypatch):
    import openpyxl
    from unittest.mock import MagicMock
    from src.tabs.skilled_aadhaar_report_tab import SkilledAadhaarReportTab

    # Mock tkinter messagebox
    monkeypatch.setattr("tkinter.messagebox.showinfo", MagicMock())
    monkeypatch.setattr("tkinter.messagebox.showerror", MagicMock())

    # Create dummy tab instance without full GUI init
    tab = SkilledAadhaarReportTab.__new__(SkilledAadhaarReportTab)
    tab.log = MagicMock()
    tab.all_scraped_data = [
        {
            "panchayat": "Burkundi",
            "sno": "1",
            "job_card_no": "JH-01-001-001/101",
            "worker_name": "RAMESH KUMAR",
            "gender": "M",
            "aadhaar_no": "XXXXXXXX1234",
            "name_as_per_aadhaar": "RAMESH KUMAR",
            "status": "Seeded"
        },
        {
            "panchayat": "Burkundi",
            "sno": "2",
            "job_card_no": "JH-01-001-001/102",
            "worker_name": "SITA DEVI",
            "gender": "F",
            "aadhaar_no": "",
            "name_as_per_aadhaar": "",
            "status": "Pending"
        },
        {
            "panchayat": "Hathidari",
            "sno": "1",
            "job_card_no": "JH-01-001-002/201",
            "worker_name": "ANIL YADAV",
            "gender": "M",
            "aadhaar_no": "XXXXXXXX5678",
            "name_as_per_aadhaar": "ANIL YADAV",
            "status": "Seeded"
        }
    ]
    tab.summary_stats = {
        "total": 3,
        "male": 2,
        "female": 1,
        "other": 0,
        "seeded": 2,
        "pending": 1
    }

    out_xlsx = str(tmp_path / "skilled_report.xlsx")
    saved_path = tab.export_professional_report(filepath=out_xlsx)

    assert saved_path == out_xlsx
    assert Path(out_xlsx).exists()

    # Verify openpyxl workbook contents & structure
    wb = openpyxl.load_workbook(out_xlsx)
    assert "Worker Details" in wb.sheetnames
    assert "Panchayat Summary" in wb.sheetnames

    ws1 = wb["Worker Details"]
    # Check title & branding
    assert "Skilled & Semi-Skilled Worker Aadhaar Status Report" in str(ws1["A1"].value)
    assert "Generated via NregaBot (https://nregabot.com)" in str(ws1["A2"].value)

    # Check KPI cards
    assert ws1["A4"].value == "👥 Total Workers"
    assert ws1["A5"].value == 3
    assert ws1["C4"].value == "👨 Male"
    assert ws1["C5"].value == 2
    assert ws1["D4"].value == "👩 Female"
    assert ws1["D5"].value == 1
    assert "Aadhaar Seeded" in str(ws1["E4"].value)
    assert ws1["E5"].value == 2
    assert "Pending" in str(ws1["G4"].value)
    assert ws1["G5"].value == 1

    # Check table headers row 7
    assert ws1.cell(row=7, column=1).value == "S.No"
    assert ws1.cell(row=7, column=2).value == "Panchayat"
    assert ws1.cell(row=7, column=8).value == "Aadhaar Status"

    # Check data rows row 8+
    assert ws1.cell(row=8, column=2).value == "Burkundi"
    assert ws1.cell(row=8, column=8).value == "✅ Seeded"
    assert ws1.cell(row=9, column=8).value == "⏳ Pending"

    # Check Sheet 2: Panchayat Summary
    ws2 = wb["Panchayat Summary"]
    assert "Panchayat-wise Skilled Worker Summary" in str(ws2["A1"].value)
    assert ws2.cell(row=4, column=2).value == "Panchayat Name"
    assert ws2.cell(row=5, column=2).value == "Burkundi"
    assert ws2.cell(row=5, column=3).value == 2  # Total for Burkundi
    assert ws2.cell(row=6, column=2).value == "Hathidari"
    assert ws2.cell(row=6, column=3).value == 1  # Total for Hathidari
    assert ws2.cell(row=7, column=2).value == "Total (सभी पंचायत)"
    assert ws2.cell(row=7, column=3).value == 3


def test_export_professional_report_csv(tmp_path, monkeypatch):
    from unittest.mock import MagicMock
    from src.tabs.skilled_aadhaar_report_tab import SkilledAadhaarReportTab

    monkeypatch.setattr("tkinter.messagebox.showinfo", MagicMock())
    monkeypatch.setattr("tkinter.messagebox.showerror", MagicMock())

    tab = SkilledAadhaarReportTab.__new__(SkilledAadhaarReportTab)
    tab.log = MagicMock()
    tab.all_scraped_data = [
        {
            "panchayat": "Burkundi",
            "sno": "1",
            "job_card_no": "JH-01-001-001/101",
            "worker_name": "RAMESH KUMAR",
            "gender": "M",
            "aadhaar_no": "XXXXXXXX1234",
            "name_as_per_aadhaar": "RAMESH KUMAR",
            "status": "Seeded"
        }
    ]
    tab.summary_stats = {"total": 1, "male": 1, "female": 0, "other": 0, "seeded": 1, "pending": 0}

    out_csv = str(tmp_path / "skilled_report.csv")
    saved_path = tab.export_professional_report(filepath=out_csv)

    assert saved_path == out_csv
    assert Path(out_csv).exists()
    content = Path(out_csv).read_text(encoding="utf-8-sig")
    assert "Burkundi" in content
    assert "RAMESH KUMAR" in content


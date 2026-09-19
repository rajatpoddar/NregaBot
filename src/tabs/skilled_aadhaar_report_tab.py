# tabs/skilled_aadhaar_report_tab.py
"""Skilled & Semi-Skilled Worker Aadhaar Updation Report Tab.

Automates the extraction of skilled/semi-skilled worker Aadhaar report from
the MGNREGA / VB-G-RAM-G portal page: SkilledBulkAadhaarUpdate.aspx.

Features:
- Panchayat selection (Single Panchayat, All Panchayats, My Saved Panchayats)
- Multi-page pagination support
- Scrapes: S.No, Job Card No, Worker Name, Gender, Aadhaar No, Name as per Aadhaar
- Computes overall & per-panchayat summary:
  * Total Workers
  * Male & Female Worker Counts
  * Aadhaar Seeded (Chada hua) vs Pending (Banki)
- Professional Excel & CSV Export
"""

import time
import threading
import json
import os
import datetime
import tkinter as tk
from tkinter import ttk, messagebox
import customtkinter as ctk

from src import config
from .base_tab import BaseAutomationTab
from src.utils import get_logger
from src.i18n import tr
from typing import Any, Dict, List, Optional, Tuple

logger = get_logger()

ALL_PANCHAYATS_LABEL = config.ALL_PANCHAYATS_LABEL
MY_PANCHAYATS_LABEL = config.MY_PANCHAYATS_LABEL


def parse_skilled_aadhaar_table(html_content: str) -> Tuple[List[Dict[str, Any]], Dict[str, int]]:
    """Parses rows from SkilledBulkAadhaarUpdate table HTML.
    
    Returns:
        tuple: (data_list, summary_dict)
    """
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html_content, "html.parser")
    table = soup.find("table", id="ctl00_ContentPlaceHolder1_Grid_debarred")
    if not table:
        return [], {"total": 0, "male": 0, "female": 0, "other": 0, "seeded": 0, "pending": 0}
    
    rows = table.find_all("tr")
    data: List[Dict[str, Any]] = []
    male = female = other = seeded = pending = 0
    
    for r in rows:
        cols = r.find_all("td")
        if len(cols) == 7:
            sno_span = cols[0].find("span", id=lambda x: x and x.endswith("_sno"))
            if not sno_span:
                continue
            
            sno = sno_span.get_text(strip=True)
            jc_span = cols[1].find("span", id=lambda x: x and x.endswith("_lbljcno"))
            jc_no = jc_span.get_text(strip=True) if jc_span else cols[1].get_text(strip=True)
            
            name_span = cols[2].find("span", id=lambda x: x and x.endswith("_lblapp_name"))
            worker_name = name_span.get_text(strip=True) if name_span else cols[2].get_text(strip=True)
            
            gender_span = cols[3].find("span", id=lambda x: x and x.endswith("_lblgender"))
            gender = gender_span.get_text(strip=True) if gender_span else cols[3].get_text(strip=True)
            
            uid_span = cols[4].find("span", id=lambda x: x and x.endswith("_lb_UID_No"))
            uid_text = uid_span.get_text(strip=True) if uid_span else cols[4].get_text(strip=True)
            
            aadhaar_name = cols[5].get_text(strip=True)
            
            has_aadhaar = bool(uid_text)
            status = "Seeded" if has_aadhaar else "Pending"
            
            gen_upper = gender.upper().strip()
            if gen_upper in ("M", "MALE"):
                male += 1
            elif gen_upper in ("F", "FEMALE"):
                female += 1
            else:
                other += 1
                
            if has_aadhaar:
                seeded += 1
            else:
                pending += 1
                
            data.append({
                "sno": sno,
                "job_card_no": jc_no,
                "worker_name": worker_name,
                "gender": gender,
                "aadhaar_no": uid_text,
                "name_as_per_aadhaar": aadhaar_name,
                "status": status
            })
            
    summary = {
        "total": len(data),
        "male": male,
        "female": female,
        "other": other,
        "seeded": seeded,
        "pending": pending
    }
    return data, summary


class SkilledAadhaarReportTab(BaseAutomationTab):
    def __init__(self, parent: Any, app_instance: Any) -> None:
        super().__init__(parent, app_instance, "skilled_aadhaar_report")
        
        if self.automation_key not in self.app.stop_events:
            self.app.stop_events[self.automation_key] = threading.Event()
        
        self.all_scraped_data: List[Dict[str, Any]] = []
        self.scraped_keys: set = set()
        
        # Summary counters
        self.summary_stats = {
            "total": 0,
            "male": 0,
            "female": 0,
            "other": 0,
            "seeded": 0,
            "pending": 0,
            "panchayats_count": 0
        }
        
        self._setup_ui()
        self.load_inputs()

    def _setup_ui(self) -> None:
        # Header card (pack-managed wrapper)
        header_wrap = ctk.CTkFrame(self, fg_color="transparent")
        header_wrap.pack(fill="x", padx=0, pady=0)
        self._create_header_card(
            header_wrap,
            "👷‍♂️",
            tr("tab.skilled_aadhaar_report.title", default="Skilled & Semi-Skilled Report"),
            tr("tab.skilled_aadhaar_report.subtitle", default="Scan skilled worker Aadhaar updation status & generate summary report."),
            icon_key="emoji_skilled_worker_report"
        )

        # TabView: Settings / Results / Logs & Status
        self.tab_view = ctk.CTkTabview(self)
        self.tab_view.pack(fill="both", expand=True, padx=10, pady=(6, 5))

        settings_tab = self.tab_view.add("Settings")
        results_tab = self.tab_view.add("Results")
        self._create_log_and_status_area(self.tab_view)

        settings_tab.grid_columnconfigure(0, weight=1)
        settings_tab.grid_rowconfigure(2, weight=1)
        results_tab.grid_columnconfigure(0, weight=1)
        results_tab.grid_rowconfigure(2, weight=1)

        # --- Settings: Input Section ---
        input_frame = ctk.CTkFrame(
            settings_tab,
            corner_radius=12,
            border_width=1,
            border_color=("gray85", "gray30"),
            fg_color=("gray97", "gray18")
        )
        input_frame.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 6))
        input_frame.grid_columnconfigure(1, weight=1)
        input_frame.grid_columnconfigure(3, weight=1)

        # Panchayat Selection Dropdown
        ctk.CTkLabel(input_frame, text="Panchayat:", font=("Arial", 12, "bold")).grid(
            row=0, column=0, padx=(12, 5), pady=12, sticky="w"
        )
        p_vals = self.app.history_manager.get_suggestions("location_panchayat") or [""]
        self.panchayat_var = ctk.StringVar(value=ALL_PANCHAYATS_LABEL)
        self.panchayat_menu = ctk.CTkOptionMenu(
            input_frame,
            variable=self.panchayat_var,
            values=self._all_panchayat_values(p_vals),
            width=200
        )
        self.panchayat_menu.grid(row=0, column=1, padx=5, pady=12, sticky="ew")

        # Filter Dropdown
        ctk.CTkLabel(input_frame, text="Filter:", font=("Arial", 12, "bold")).grid(
            row=0, column=2, padx=(15, 5), pady=12, sticky="w"
        )
        self.filter_var = ctk.StringVar(value="All")
        self.filter_var.trace_add("write", lambda *_: self.apply_filter_visuals())
        self.filter_menu = ctk.CTkOptionMenu(
            input_frame,
            variable=self.filter_var,
            values=["All", "Aadhaar Seeded", "Aadhaar Pending"],
            width=160
        )
        self.filter_menu.grid(row=0, column=3, padx=(5, 12), pady=12, sticky="ew")

        note_label = ctk.CTkLabel(
            settings_tab,
            text="💡 Note: Select '🌐 All Panchayats' to scan all panchayats, or '⭐ My Saved Panchayats' for your saved panchayats.",
            text_color=("gray40", "gray70"),
            font=("Arial", 11, "italic")
        )
        note_label.grid(row=1, column=0, sticky="w", padx=20, pady=(4, 4))

        # --- Settings: Stats Frame ---
        self.stats_frame = ctk.CTkFrame(
            settings_tab,
            corner_radius=12,
            border_width=1,
            border_color=("gray85", "gray30"),
            fg_color=("gray97", "gray18")
        )
        self.stats_frame.grid(row=2, column=0, sticky="nsew", padx=12, pady=(0, 5))
        self.stats_frame.grid_columnconfigure(0, weight=1)
        
        self.stats_label_header = ctk.CTkLabel(
            self.stats_frame,
            text="📊 Scan Summary / सारांश:",
            font=("Arial", 12, "bold")
        )
        self.stats_label_header.pack(anchor="w", padx=12, pady=(8, 4))
        
        self.stats_text = ctk.CTkLabel(
            self.stats_frame,
            text="(No data scanned yet. Click 'Start Automation' to begin.)",
            font=("Arial", 11),
            justify="left",
            wraplength=850
        )
        self.stats_text.pack(anchor="w", padx=16, pady=(0, 8))

        # --- Settings: Action Buttons ---
        self._create_action_buttons(parent_frame=settings_tab).grid(row=3, column=0, pady=(8, 6))

        # --- Results Tab UI ---
        # Top summary bar in Results tab
        self.results_summary_frame = ctk.CTkFrame(
            results_tab,
            corner_radius=10,
            fg_color=("gray90", "gray22"),
            border_width=1,
            border_color=("gray80", "gray35")
        )
        self.results_summary_frame.grid(row=0, column=0, sticky="ew", padx=10, pady=(8, 6))
        self.results_summary_frame.grid_columnconfigure((0, 1, 2, 3, 4), weight=1)

        # Summary Metric Badges
        self.lbl_card_total = ctk.CTkLabel(
            self.results_summary_frame,
            text="Total Workers\n0",
            font=("Arial", 12, "bold"),
            text_color=("gray10", "gray90")
        )
        self.lbl_card_total.grid(row=0, column=0, padx=8, pady=8)

        self.lbl_card_gender = ctk.CTkLabel(
            self.results_summary_frame,
            text="Male / Female\n0 / 0",
            font=("Arial", 12, "bold"),
            text_color=("blue", "#4A90E2")
        )
        self.lbl_card_gender.grid(row=0, column=1, padx=8, pady=8)

        self.lbl_card_seeded = ctk.CTkLabel(
            self.results_summary_frame,
            text="Aadhaar Seeded (चढ़ा हुआ)\n0",
            font=("Arial", 12, "bold"),
            text_color=("darkgreen", "#2ECC71")
        )
        self.lbl_card_seeded.grid(row=0, column=2, padx=8, pady=8)

        self.lbl_card_pending = ctk.CTkLabel(
            self.results_summary_frame,
            text="Aadhaar Pending (बाकी)\n0",
            font=("Arial", 12, "bold"),
            text_color=("darkred", "#E74C3C")
        )
        self.lbl_card_pending.grid(row=0, column=3, padx=8, pady=8)

        # Export Button in Results tab
        export_frame = ctk.CTkFrame(results_tab, fg_color="transparent")
        export_frame.grid(row=1, column=0, sticky="ew", padx=10, pady=(0, 6))
        
        self.export_btn = ctk.CTkButton(
            export_frame,
            text="📥 Export to Excel / CSV",
            command=self.export_professional_report,
            state="disabled",
            fg_color=config.COLORS.get("green_export", "#27ae60"),
            font=("Arial", 12, "bold")
        )
        self.export_btn.pack(side="left", padx=5)

        # Treeview Data Table
        table_container = ctk.CTkFrame(results_tab, fg_color="transparent")
        table_container.grid(row=2, column=0, sticky="nsew", padx=10, pady=(0, 8))
        table_container.grid_rowconfigure(0, weight=1)
        table_container.grid_columnconfigure(0, weight=1)

        columns = ("Panchayat", "SNo", "JobCardNo", "WorkerName", "Gender", "AadhaarNo", "NameAsPerAadhaar", "AadhaarStatus")
        self.tree = ttk.Treeview(table_container, columns=columns, show="headings", selectmode="browse")
        
        headers = {
            "Panchayat": "Panchayat",
            "SNo": "S.No",
            "JobCardNo": "Job Card Number",
            "WorkerName": "Worker Name",
            "Gender": "Gender",
            "AadhaarNo": "Aadhaar Number",
            "NameAsPerAadhaar": "Name as per Aadhaar",
            "AadhaarStatus": "Aadhaar Status"
        }
        widths = {
            "Panchayat": 120,
            "SNo": 50,
            "JobCardNo": 180,
            "WorkerName": 160,
            "Gender": 60,
            "AadhaarNo": 140,
            "NameAsPerAadhaar": 160,
            "AadhaarStatus": 110
        }
        
        for col, h in headers.items():
            self.tree.heading(col, text=h, anchor="center")
            self.tree.column(col, width=widths[col], anchor="center" if col in ("SNo", "Gender", "AadhaarStatus") else "w")

        vsb = ttk.Scrollbar(table_container, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(table_container, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

    # --- UI Helper & Event Handlers ---
    def apply_filter_visuals(self) -> None:
        """Applies visual filter on the results Treeview."""
        filter_val = self.filter_var.get()
        self.tree.delete(*self.tree.get_children())
        
        for item in self.all_scraped_data:
            st = item.get("status", "")
            if filter_val == "Aadhaar Seeded" and st != "Seeded":
                continue
            if filter_val == "Aadhaar Pending" and st != "Pending":
                continue
                
            self.tree.insert("", "end", values=(
                item.get("panchayat", ""),
                item.get("sno", ""),
                item.get("job_card_no", ""),
                item.get("worker_name", ""),
                item.get("gender", ""),
                item.get("aadhaar_no", ""),
                item.get("name_as_per_aadhaar", ""),
                item.get("status", "")
            ))

    def _update_summary_cards(self) -> None:
        """Updates top summary badges and text."""
        tot = self.summary_stats["total"]
        m = self.summary_stats["male"]
        f = self.summary_stats["female"]
        o = self.summary_stats["other"]
        s = self.summary_stats["seeded"]
        p = self.summary_stats["pending"]

        self.lbl_card_total.configure(text=f"Total Workers\n{tot}")
        self.lbl_card_gender.configure(text=f"Male / Female / Other\n{m} / {f} / {o}")
        self.lbl_card_seeded.configure(text=f"Aadhaar Seeded (चढ़ा हुआ)\n{s}")
        self.lbl_card_pending.configure(text=f"Aadhaar Pending (बाकी)\n{p}")

        text_summary = (
            f"• Total Scraped Workers: {tot}\n"
            f"• Male Workers: {m} | Female Workers: {f} | Other: {o}\n"
            f"• Aadhaar Number Seeded (चढ़ा हुआ): {s}\n"
            f"• Aadhaar Number Pending (बाकी): {p}"
        )
        self.stats_text.configure(text=text_summary)

        if tot > 0:
            self.export_btn.configure(state="normal")

    def log(self, message: str, level: str = "INFO") -> None:
        """Helper to log messages to tab's log_display text widget."""
        lvl = str(level).upper()
        log_disp = getattr(self, "log_display", None)
        if lvl in ("ERROR", "FAILED"):
            self.app.log_message(log_disp, message, "error")
        elif lvl in ("WARNING", "WARN"):
            self.app.log_message(log_disp, message, "warning")
        elif lvl in ("SUCCESS", "DONE"):
            self.app.log_message(log_disp, message, "success")
        else:
            self.app.log_message(log_disp, message, "info")

    # --- History & Input Saving ---
    def save_inputs(self) -> None:
        data = {
            "panchayat": self.panchayat_var.get(),
            "filter": self.filter_var.get()
        }
        self.app.history_manager.save_tab_inputs_batch("skilled_aadhaar_report", data)

    def load_inputs(self) -> None:
        data = self.app.history_manager.get_tab_inputs("skilled_aadhaar_report")
        if data:
            if "panchayat" in data:
                self.panchayat_var.set(data["panchayat"])
            if "filter" in data:
                self.filter_var.set(data["filter"])

    def start_automation(self) -> None:
        """Starts the automation workflow in a background thread."""
        self.save_inputs()
        self.set_common_ui_state(running=True)
        self.export_btn.configure(state="disabled")
        
        self.all_scraped_data.clear()
        self.scraped_keys.clear()
        self.summary_stats = {"total": 0, "male": 0, "female": 0, "other": 0, "seeded": 0, "pending": 0, "panchayats_count": 0}
        
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.stats_text.configure(text="(Processing...)")

        self.app.start_automation_thread(
            self.automation_key,
            self._run_process_safe,
            args=()
        )

    def _run_process_safe(self) -> None:
        """Called by automation thread runner with active driver."""
        try:
            driver = None
            if hasattr(self.app, "browser_manager"):
                driver = self.app.browser_manager.get_driver()
            if not driver and hasattr(self.app, "get_driver_or_launch"):
                driver = self.app.get_driver_or_launch()
            if not driver and hasattr(self.app, "get_driver"):
                driver = self.app.get_driver()
            if not driver and hasattr(self.app, "driver"):
                driver = self.app.driver

            if not driver:
                self.log("❌ Driver not available. Please launch browser first.", "ERROR")
                return
            self.run_process(driver)
        except Exception as e:
            self.log(f"❌ Automation encountered an error: {e}", "ERROR")
            logger.exception(f"Error in SkilledAadhaarReportTab: {e}")
            self.handle_error(e)
        finally:
            self.app.after(0, self.set_common_ui_state, False)

    # --- Main Automation Method ---
    def run_process(self, driver: Any) -> None:
        """Runs the automation scraping thread."""
        from ._imports import By, Select, WebDriverWait, EC, NoSuchElementException, TimeoutException

        self.save_inputs()
        stop_event = self.app.stop_events[self.automation_key]
        
        # Reset data
        self.all_scraped_data.clear()
        self.scraped_keys.clear()
        self.summary_stats = {"total": 0, "male": 0, "female": 0, "other": 0, "seeded": 0, "pending": 0, "panchayats_count": 0}

        def _safe_ui(cb):
            self.app.after(0, cb)

        _safe_ui(lambda: [
            self.tree.delete(*self.tree.get_children()),
            self.export_btn.configure(state="disabled"),
            self._update_summary_cards()
        ])

        self.log(tr("msg.starting_automation", default="Starting Skilled & Semi-Skilled Aadhaar Report Scraper..."), "INFO")

        selected_panchayat_input = self.panchayat_var.get().strip()
        
        # Ensure driver is on the correct page
        target_url = self.resolve_portal_url(config.SKILLED_AADHAAR_CONFIG["url"])
        curr_url = driver.current_url or ""
        if "skilledbulkaadhaarupdate.aspx" not in curr_url.lower():
            self.log(f"🌐 Navigating to Skilled/Semi-Skilled page ({target_url})...", "INFO")
            try:
                driver.get(target_url)
                time.sleep(3)
            except Exception as ex:
                self.log(f"❌ Failed to open URL {target_url}: {ex}", "ERROR")
                return

        wait = WebDriverWait(driver, 15)
        panch_dropdown_id = "ctl00_ContentPlaceHolder1_ddl_panch"
        is_gp_login = False

        # Detect whether we have a Panchayat dropdown (Block/PO login) or GP login
        dropdown_elem = None
        try:
            dropdown_elem = wait.until(EC.presence_of_element_located((By.ID, panch_dropdown_id)))
        except TimeoutException:
            if "skilledbulkaadhaarupdate.aspx" in (driver.current_url or "").lower():
                is_gp_login = True
                self.log("ℹ️ Panchayat dropdown not found — operating in GP login mode.", "INFO")
            else:
                self.log(f"❌ Could not find Panchayat dropdown on page ({driver.current_url}). Please ensure you are logged into the portal.", "ERROR")
                return

        # Determine target Panchayats list
        target_panchayats: List[str] = []
        if is_gp_login:
            gp_name = (
                self._clean_panchayat_value(selected_panchayat_input)
                or self._read_gp_panchayat(driver, ["ctl00_ContentPlaceHolder1_panchlbl", "ctl00_ContentPlaceHolder1_lbl_panch"])
                or "Gram Panchayat"
            )
            target_panchayats = [gp_name]
        else:
            select_obj = Select(dropdown_elem)
            available_options = [
                opt.text.strip() for opt in select_obj.options 
                if opt.text.strip() and not self._is_aggregate_panchayat_name(opt.text) and "select" not in opt.text.lower()
            ]

            if self._is_panchayat_label(selected_panchayat_input):
                if self._is_my_saved_panchayat(selected_panchayat_input):
                    saved = self._get_saved_panchayats()
                    target_panchayats = [p for p in available_options if any(p.lower() == s.lower() for s in saved)]
                    if not target_panchayats:
                        self.log("⚠️ No matching saved panchayats found in dropdown.", "WARNING")
                        return
                    self.log(f"⭐ My Saved Panchayats mode: {len(target_panchayats)} saved panchayat(s) will be processed.", "INFO")
                else:
                    # 🌐 All Panchayats
                    target_panchayats = available_options
                    self.log(f"🌐 All Panchayats mode: found {len(target_panchayats)} panchayats.", "INFO")
            else:
                # Single panchayat
                clean_input = self._clean_panchayat_value(selected_panchayat_input) or selected_panchayat_input
                matched = [p for p in available_options if p.lower() == clean_input.lower()]
                if not matched:
                    matched = [p for p in available_options if clean_input.lower() in p.lower()]
                if matched:
                    target_panchayats = [matched[0]]
                else:
                    target_panchayats = [clean_input]

        if not target_panchayats:
            self.log("⚠️ No panchayats to process.", "WARNING")
            return

        preview_str = ", ".join(target_panchayats[:5])
        if len(target_panchayats) > 5:
            preview_str += f"... (+{len(target_panchayats)-5} more)"
        self.log(f"📋 Processing {len(target_panchayats)} Panchayat(s): {preview_str}", "INFO")

        for idx, pan_name in enumerate(target_panchayats, start=1):
            if stop_event.is_set():
                self.log("⏹️ Stopped by user.", "WARNING")
                break

            self.log(f"[{idx}/{len(target_panchayats)}] Selecting Panchayat: {pan_name}", "INFO")
            self.update_status(f"Scanning {pan_name} ({idx}/{len(target_panchayats)})", idx / max(len(target_panchayats), 1))

            if not is_gp_login:
                try:
                    dropdown_elem = wait.until(EC.element_to_be_clickable((By.ID, panch_dropdown_id)))
                    select_obj = Select(dropdown_elem)

                    curr_selected_text = ""
                    try:
                        curr_selected_text = select_obj.first_selected_option.text.strip()
                    except Exception:
                        pass

                    # Match option
                    matched_opt_text = None
                    for opt in select_obj.options:
                        opt_txt = opt.text.strip()
                        if opt_txt.lower() == pan_name.lower():
                            matched_opt_text = opt.text
                            break
                    if not matched_opt_text:
                        for opt in select_obj.options:
                            opt_txt = opt.text.strip()
                            if pan_name.lower() in opt_txt.lower():
                                matched_opt_text = opt.text
                                break

                    if not matched_opt_text:
                        self.log(f"⚠️ Panchayat '{pan_name}' not found in dropdown options.", "WARNING")
                        continue

                    # Check if already selected with table present AND at first panchayat
                    grid_elements = driver.find_elements(By.ID, "ctl00_ContentPlaceHolder1_Grid_debarred")
                    if curr_selected_text.lower() == matched_opt_text.lower() and grid_elements and idx == 1:
                        self.log(f"  └ Panchayat '{matched_opt_text}' already selected.", "INFO")
                    else:
                        old_grid = grid_elements[0] if grid_elements else None
                        select_obj.select_by_visible_text(matched_opt_text)

                        # Trigger ASP.NET postback
                        try:
                            driver.execute_script("__doPostBack('ctl00$ContentPlaceHolder1$ddl_panch', '');")
                        except Exception:
                            driver.execute_script("arguments[0].dispatchEvent(new Event('change', { bubbles: true }));", dropdown_elem)

                        # Wait for page reload to complete
                        if old_grid:
                            try:
                                WebDriverWait(driver, 10).until(EC.staleness_of(old_grid))
                            except Exception:
                                time.sleep(2)
                        else:
                            try:
                                WebDriverWait(driver, 10).until(EC.staleness_of(dropdown_elem))
                            except Exception:
                                time.sleep(2)

                        wait.until(EC.presence_of_element_located((By.ID, panch_dropdown_id)))
                        time.sleep(1.5)

                    # Ensure Grid is on Page 1 (ASP.NET GridView retains PageIndex across dropdown changes)
                    try:
                        active_spans = driver.find_elements(
                            By.XPATH,
                            "//table[@id='ctl00_ContentPlaceHolder1_Grid_debarred']//table//span"
                        )
                        if active_spans:
                            active_p = active_spans[0].text.strip()
                            if active_p and active_p != "1":
                                self.log(f"  └ Grid loaded at Page {active_p}; resetting to Page 1 for {pan_name}...", "INFO")
                                current_grid = driver.find_element(By.ID, "ctl00_ContentPlaceHolder1_Grid_debarred")
                                driver.execute_script("__doPostBack('ctl00$ContentPlaceHolder1$Grid_debarred', 'Page$1');")
                                try:
                                    WebDriverWait(driver, 10).until(EC.staleness_of(current_grid))
                                except Exception:
                                    time.sleep(2)
                                wait.until(EC.presence_of_element_located((By.ID, "ctl00_ContentPlaceHolder1_Grid_debarred")))
                                time.sleep(1.5)
                    except Exception as ex:
                        logger.debug(f"Reset to page 1 check error: {ex}")

                except Exception as ex:
                    self.log(f"❌ Error selecting Panchayat {pan_name}: {ex}", "ERROR")
                    continue

            # Scrape pages for this Panchayat
            visited_pages = set()
            current_page_num = 1
            while not stop_event.is_set():
                # Detect actual page number from DOM pager if present
                try:
                    active_spans = driver.find_elements(
                        By.XPATH,
                        "//table[@id='ctl00_ContentPlaceHolder1_Grid_debarred']//table//span"
                    )
                    if active_spans and active_spans[0].text.strip().isdigit():
                        current_page_num = int(active_spans[0].text.strip())
                except Exception:
                    pass

                if current_page_num in visited_pages:
                    self.log(f"  └ Page {current_page_num} already processed for {pan_name}. Finishing Panchayat.", "INFO")
                    break
                visited_pages.add(current_page_num)

                self.log(f"  └ Scanning Page {current_page_num} for {pan_name}...", "INFO")

                page_data, _ = parse_skilled_aadhaar_table(driver.page_source)

                # Fallback: if page 1 returned 0 records, try Page$1 reset in case ViewState PageIndex was out-of-bounds
                if not page_data and len(visited_pages) == 1:
                    try:
                        grid_to_check = driver.find_element(By.ID, "ctl00_ContentPlaceHolder1_Grid_debarred")
                        driver.execute_script("__doPostBack('ctl00$ContentPlaceHolder1$Grid_debarred', 'Page$1');")
                        try:
                            WebDriverWait(driver, 8).until(EC.staleness_of(grid_to_check))
                        except Exception:
                            time.sleep(2)
                        wait.until(EC.presence_of_element_located((By.ID, "ctl00_ContentPlaceHolder1_Grid_debarred")))
                        time.sleep(1.5)
                        page_data, _ = parse_skilled_aadhaar_table(driver.page_source)
                    except Exception:
                        pass

                if not page_data:
                    if len(visited_pages) == 1:
                        self.log(f"  └ No worker records found for {pan_name}.", "INFO")
                    break

                page_rows_added = 0
                for item in page_data:
                    jc_no = item["job_card_no"]
                    w_name = item["worker_name"]
                    row_key = f"{pan_name}_{jc_no}_{w_name}"
                    if row_key in self.scraped_keys:
                        continue

                    self.scraped_keys.add(row_key)
                    item_with_pan = dict(item)
                    item_with_pan["panchayat"] = pan_name
                    self.all_scraped_data.append(item_with_pan)

                    gen_upper = item["gender"].upper().strip()
                    if gen_upper in ("M", "MALE"):
                        self.summary_stats["male"] += 1
                    elif gen_upper in ("F", "FEMALE"):
                        self.summary_stats["female"] += 1
                    else:
                        self.summary_stats["other"] += 1

                    if item["status"] == "Seeded":
                        self.summary_stats["seeded"] += 1
                    else:
                        self.summary_stats["pending"] += 1

                    self.summary_stats["total"] += 1
                    page_rows_added += 1

                self.log(f"  └ Page {current_page_num}: {page_rows_added} new worker records recorded.", "INFO")

                _safe_ui(lambda: [
                    self.apply_filter_visuals(),
                    self._update_summary_cards()
                ])

                # Check for next page link
                next_page_num = current_page_num + 1
                next_link = None
                try:
                    candidates = driver.find_elements(
                        By.XPATH,
                        f"//table[@id='ctl00_ContentPlaceHolder1_Grid_debarred']//table//a[contains(@href, 'Page${next_page_num}') or (normalize-space()='{next_page_num}' and contains(@href, '__doPostBack'))]"
                    )
                    if candidates:
                        next_link = candidates[0]
                except Exception as ex:
                    logger.debug(f"Pagination check: {ex}")
                    next_link = None

                if next_link:
                    try:
                        self.log(f"  └ Navigating to Page {next_page_num} for {pan_name}...", "INFO")
                        grid_to_wait = driver.find_element(By.ID, "ctl00_ContentPlaceHolder1_Grid_debarred")
                        driver.execute_script(f"__doPostBack('ctl00$ContentPlaceHolder1$Grid_debarred', 'Page${next_page_num}');")
                        try:
                            WebDriverWait(driver, 10).until(EC.staleness_of(grid_to_wait))
                        except Exception:
                            time.sleep(2)
                        wait.until(EC.presence_of_element_located((By.ID, "ctl00_ContentPlaceHolder1_Grid_debarred")))
                        time.sleep(1.5)
                    except Exception as ex:
                        self.log(f"  └ Could not navigate to page {next_page_num}: {ex}", "WARNING")
                        break
                else:
                    # No more pages for this Panchayat
                    break

        self.log(f"✅ Scraping finished! Scraped {self.summary_stats['total']} worker records across {len(target_panchayats)} Panchayat(s).", "SUCCESS")
        self.update_status("Finished", 1.0)
        self.activity_details = f"Scraped {self.summary_stats['total']} worker records across {len(target_panchayats)} Panchayat(s)."
        self.show_automation_notification("success")
        _safe_ui(self._update_summary_cards)

    # --- Export Report ---
    def export_professional_report(self, filepath: Optional[str] = None, auto_open: bool = True) -> Optional[str]:
        """Exports scraped data and summary to an executive-grade Excel workbook or CSV."""
        if not self.all_scraped_data:
            messagebox.showinfo("Export", "No data to export.")
            return None

        if not filepath:
            from tkinter import filedialog
            default_filename = f"Skilled_Aadhaar_Report_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            reports_dir = self.app.get_report_path(self._report_category())
            os.makedirs(reports_dir, exist_ok=True)
            filepath = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel Workbook", "*.xlsx"), ("CSV Files", "*.csv")],
                initialdir=reports_dir,
                initialfile=default_filename,
                title="Save Skilled Worker Report"
            )

        if not filepath:
            return None

        try:
            if filepath.endswith(".csv"):
                import csv
                with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    writer.writerow(["Panchayat", "S.No", "Job Card No", "Worker Name", "Gender", "Aadhaar No", "Name as per Aadhaar", "Aadhaar Status"])
                    for r in self.all_scraped_data:
                        writer.writerow([
                            r.get("panchayat", ""),
                            r.get("sno", ""),
                            r.get("job_card_no", ""),
                            r.get("worker_name", ""),
                            r.get("gender", ""),
                            r.get("aadhaar_no", ""),
                            r.get("name_as_per_aadhaar", ""),
                            r.get("status", "")
                        ])
            else:
                import sys
                import openpyxl
                from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
                from openpyxl.utils import get_column_letter

                # Recalculate summary metrics if needed
                total = self.summary_stats.get("total", 0) or len(self.all_scraped_data)
                male = self.summary_stats.get("male", 0)
                female = self.summary_stats.get("female", 0)
                other = self.summary_stats.get("other", 0)
                seeded = self.summary_stats.get("seeded", 0)
                pending = self.summary_stats.get("pending", 0)

                if total > 0 and (male + female + other == 0):
                    for r in self.all_scraped_data:
                        g = str(r.get("gender", "")).upper()
                        if g in ("M", "MALE"):
                            male += 1
                        elif g in ("F", "FEMALE"):
                            female += 1
                        else:
                            other += 1
                        st = str(r.get("status", "")).lower()
                        if "seed" in st or "chada" in st or "चढ़ा" in st:
                            seeded += 1
                        else:
                            pending += 1

                seeded_ratio = (seeded / total * 100) if total else 0.0

                wb = openpyxl.Workbook()

                # Reusable styles
                navy_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
                sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
                even_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
                odd_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
                total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

                title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
                sub_font = Font(name="Calibri", size=9, italic=True, color="475569")
                hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                bold_font = Font(name="Calibri", size=10, bold=True, color="1E293B")
                regular_font = Font(name="Calibri", size=10, color="1E293B")

                c_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

                border_thin = Border(
                    left=Side(style="thin", color="CBD5E1"),
                    right=Side(style="thin", color="CBD5E1"),
                    top=Side(style="thin", color="CBD5E1"),
                    bottom=Side(style="thin", color="CBD5E1")
                )
                border_total = Border(
                    left=Side(style="thin", color="CBD5E1"),
                    right=Side(style="thin", color="CBD5E1"),
                    top=Side(style="thin", color="64748B"),
                    bottom=Side(style="double", color="1F4E79")
                )

                def style_box(sheet, range_str, font, fill, align=c_center, border=border_thin):
                    for row_cells in sheet[range_str]:
                        for cell in row_cells:
                            if font:
                                cell.font = font
                            if fill:
                                cell.fill = fill
                            if align:
                                cell.alignment = align
                            if border:
                                cell.border = border

                # ════════════════════════════════════════════════════════════
                # SHEET 1: Worker Details
                # ════════════════════════════════════════════════════════════
                ws1 = wb.active
                ws1.title = "Worker Details"
                ws1.views.sheetView[0].showGridLines = True

                # Row 1: Title Banner
                ws1.merge_cells("A1:H1")
                ws1["A1"] = "MGNREGA / VB-G-RAM-G — Skilled & Semi-Skilled Worker Aadhaar Status Report"
                ws1.row_dimensions[1].height = 36
                style_box(ws1, "A1:H1", title_font, navy_fill)

                # Row 2: Subtitle Branding Banner
                ws1.merge_cells("A2:H2")
                now_str = datetime.datetime.now().strftime('%d-%b-%Y %I:%M %p')
                ws1["A2"] = f"Generated via NregaBot (https://nregabot.com) | Date: {now_str}"
                ws1.row_dimensions[2].height = 22
                style_box(ws1, "A2:H2", sub_font, sub_fill)

                ws1.row_dimensions[3].height = 10

                # Rows 4-5: Executive KPI Metric Cards across 8 columns
                # Card 1: Total Workers (A-B)
                ws1.merge_cells("A4:B4")
                ws1["A4"] = "👥 Total Workers"
                style_box(ws1, "A4:B4", Font(name="Calibri", size=9, bold=True, color="1F4E79"),
                          PatternFill(start_color="DCE6F1", fill_type="solid"))

                ws1.merge_cells("A5:B5")
                ws1["A5"] = total
                style_box(ws1, "A5:B5", Font(name="Calibri", size=13, bold=True, color="1F4E79"),
                          PatternFill(start_color="F2F6FA", fill_type="solid"))

                # Card 2: Male Workers (C)
                ws1["C4"] = "👨 Male"
                style_box(ws1, "C4:C4", Font(name="Calibri", size=9, bold=True, color="1D4ED8"),
                          PatternFill(start_color="DBEAFE", fill_type="solid"))
                ws1["C5"] = male
                style_box(ws1, "C5:C5", Font(name="Calibri", size=13, bold=True, color="1D4ED8"),
                          PatternFill(start_color="EFF6FF", fill_type="solid"))

                # Card 3: Female Workers (D)
                ws1["D4"] = "👩 Female"
                style_box(ws1, "D4:D4", Font(name="Calibri", size=9, bold=True, color="9D174D"),
                          PatternFill(start_color="FCE7F3", fill_type="solid"))
                ws1["D5"] = female
                style_box(ws1, "D5:D5", Font(name="Calibri", size=13, bold=True, color="9D174D"),
                          PatternFill(start_color="FDF2F8", fill_type="solid"))

                # Card 4: Aadhaar Seeded (E-F)
                ws1.merge_cells("E4:F4")
                ws1["E4"] = "✅ Aadhaar Seeded (चढ़ा हुआ)"
                style_box(ws1, "E4:F4", Font(name="Calibri", size=9, bold=True, color="166534"),
                          PatternFill(start_color="DCFCE7", fill_type="solid"))

                ws1.merge_cells("E5:F5")
                ws1["E5"] = seeded
                style_box(ws1, "E5:F5", Font(name="Calibri", size=13, bold=True, color="166534"),
                          PatternFill(start_color="F0FDF4", fill_type="solid"))

                # Card 5: Aadhaar Pending (G)
                ws1["G4"] = "⏳ Pending (बाकी)"
                style_box(ws1, "G4:G4", Font(name="Calibri", size=9, bold=True, color="991B1B"),
                          PatternFill(start_color="FEE2E2", fill_type="solid"))
                ws1["G5"] = pending
                style_box(ws1, "G5:G5", Font(name="Calibri", size=13, bold=True, color="991B1B"),
                          PatternFill(start_color="FEF2F2", fill_type="solid"))

                # Card 6: Seeded Ratio % (H)
                ws1["H4"] = "📊 Seeded %"
                style_box(ws1, "H4:H4", Font(name="Calibri", size=9, bold=True, color="155E75"),
                          PatternFill(start_color="CFFAFE", fill_type="solid"))
                ws1["H5"] = f"{seeded_ratio:.1f}%"
                style_box(ws1, "H5:H5", Font(name="Calibri", size=13, bold=True, color="155E75"),
                          PatternFill(start_color="ECFEFF", fill_type="solid"))

                ws1.row_dimensions[4].height = 20
                ws1.row_dimensions[5].height = 26
                ws1.row_dimensions[6].height = 12

                # Row 7: Data Table Headers
                headers = [
                    "S.No", "Panchayat", "Job Card No", "Worker Name",
                    "Gender", "Aadhaar No", "Name as per Aadhaar", "Aadhaar Status"
                ]
                ws1.row_dimensions[7].height = 26
                for col_idx, h in enumerate(headers, 1):
                    cell = ws1.cell(row=7, column=col_idx, value=h)
                    cell.font = hdr_font
                    cell.fill = navy_fill
                    cell.alignment = c_center
                    cell.border = border_thin

                # Row 8+: Data Rows
                green_fill = PatternFill(start_color="DCFCE7", fill_type="solid")
                green_font = Font(name="Calibri", size=10, bold=True, color="166534")
                red_fill = PatternFill(start_color="FEE2E2", fill_type="solid")
                red_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

                for idx, r in enumerate(self.all_scraped_data):
                    row_num = 8 + idx
                    ws1.row_dimensions[row_num].height = 21
                    is_even = (idx % 2 == 0)
                    row_fill = even_fill if is_even else odd_fill

                    sno = r.get("sno") or str(idx + 1)
                    pan = r.get("panchayat", "")
                    jc = r.get("job_card_no", "")
                    w_name = r.get("worker_name", "")
                    gender = r.get("gender", "")
                    aadhaar = r.get("aadhaar_no", "")
                    aadhaar_name = r.get("name_as_per_aadhaar", "")
                    status_raw = r.get("status", "")

                    row_vals = [sno, pan, jc, w_name, gender, aadhaar, aadhaar_name]
                    for col_idx, val in enumerate(row_vals, 1):
                        cell = ws1.cell(row=row_num, column=col_idx, value=str(val) if val is not None else "")
                        cell.fill = row_fill
                        cell.font = regular_font
                        cell.border = border_thin
                        cell.alignment = c_center if col_idx in (1, 3, 5, 6) else c_left

                    # Status Cell (Column 8)
                    st_cell = ws1.cell(row=row_num, column=8)
                    st_lower = str(status_raw).lower()
                    if "seed" in st_lower or "chada" in st_lower or "चढ़ा" in st_lower:
                        st_cell.value = "✅ Seeded"
                        st_cell.fill = green_fill
                        st_cell.font = green_font
                    else:
                        st_cell.value = "⏳ Pending"
                        st_cell.fill = red_fill
                        st_cell.font = red_font
                    st_cell.alignment = c_center
                    st_cell.border = border_thin

                # Freeze panes & column widths for Sheet 1
                ws1.freeze_panes = "A8"
                for col in range(1, 9):
                    max_len = len(headers[col - 1]) + 2
                    for row_idx in range(8, min(8 + len(self.all_scraped_data), 108)):
                        cv = ws1.cell(row=row_idx, column=col).value or ""
                        max_len = max(max_len, int(len(str(cv)) * 1.15))
                    col_letter = get_column_letter(col)
                    ws1.column_dimensions[col_letter].width = max(min(max_len, 42), 10)

                self._apply_print_setup(ws1, header_row=7)

                # ════════════════════════════════════════════════════════════
                # SHEET 2: Panchayat Summary Breakdown
                # ════════════════════════════════════════════════════════════
                ws2 = wb.create_sheet(title="Panchayat Summary")
                ws2.views.sheetView[0].showGridLines = True

                # Row 1: Title Banner
                ws2.merge_cells("A1:H1")
                ws2["A1"] = "MGNREGA / VB-G-RAM-G — Panchayat-wise Skilled Worker Summary"
                ws2.row_dimensions[1].height = 36
                style_box(ws2, "A1:H1", title_font, navy_fill)

                # Row 2: Subtitle Branding Banner
                ws2.merge_cells("A2:H2")
                ws2["A2"] = f"Generated via NregaBot (https://nregabot.com) | Date: {now_str}"
                ws2.row_dimensions[2].height = 22
                style_box(ws2, "A2:H2", sub_font, sub_fill)

                ws2.row_dimensions[3].height = 12

                # Row 4: Summary Headers
                sum_headers = [
                    "S.No", "Panchayat Name", "Total Workers", "Male", "Female",
                    "Aadhaar Seeded", "Aadhaar Pending", "Seeded %"
                ]
                ws2.row_dimensions[4].height = 26
                for col_idx, h in enumerate(sum_headers, 1):
                    cell = ws2.cell(row=4, column=col_idx, value=h)
                    cell.font = hdr_font
                    cell.fill = navy_fill
                    cell.alignment = c_center
                    cell.border = border_thin

                # Group by Panchayat
                pan_groups: Dict[str, Dict[str, int]] = {}
                for r in self.all_scraped_data:
                    p_name = r.get("panchayat") or "Unknown"
                    if p_name not in pan_groups:
                        pan_groups[p_name] = {"total": 0, "male": 0, "female": 0, "seeded": 0, "pending": 0}
                    pan_groups[p_name]["total"] += 1
                    g = str(r.get("gender", "")).upper()
                    if g in ("M", "MALE"):
                        pan_groups[p_name]["male"] += 1
                    elif g in ("F", "FEMALE"):
                        pan_groups[p_name]["female"] += 1

                    st = str(r.get("status", "")).lower()
                    if "seed" in st or "chada" in st or "चढ़ा" in st:
                        pan_groups[p_name]["seeded"] += 1
                    else:
                        pan_groups[p_name]["pending"] += 1

                # Write Panchayat rows
                curr_row = 5
                for p_idx, (p_name, p_stats) in enumerate(pan_groups.items(), 1):
                    ws2.row_dimensions[curr_row].height = 22
                    is_even = (p_idx % 2 == 0)
                    r_fill = even_fill if is_even else odd_fill

                    p_tot = p_stats["total"]
                    p_seed = p_stats["seeded"]
                    p_pct = (p_seed / p_tot * 100) if p_tot else 0.0

                    row_vals = [
                        p_idx, p_name, p_tot, p_stats["male"], p_stats["female"],
                        p_seed, p_stats["pending"], f"{p_pct:.1f}%"
                    ]

                    for col_idx, val in enumerate(row_vals, 1):
                        cell = ws2.cell(row=curr_row, column=col_idx, value=val)
                        cell.fill = r_fill
                        cell.border = border_thin
                        if col_idx == 2:
                            cell.font = regular_font
                            cell.alignment = c_left
                        elif col_idx == 6:
                            cell.font = green_font
                            cell.alignment = c_center
                        elif col_idx == 7:
                            cell.font = red_font
                            cell.alignment = c_center
                        elif col_idx == 8:
                            cell.font = Font(name="Calibri", size=10, bold=True, color="155E75")
                            cell.alignment = c_center
                        else:
                            cell.font = regular_font
                            cell.alignment = c_center

                    curr_row += 1

                # Grand Total Row
                ws2.row_dimensions[curr_row].height = 26
                tot_vals = [
                    "", "Total (सभी पंचायत)", total, male, female,
                    seeded, pending, f"{seeded_ratio:.1f}%"
                ]
                for col_idx, val in enumerate(tot_vals, 1):
                    cell = ws2.cell(row=curr_row, column=col_idx, value=val)
                    cell.fill = total_fill
                    cell.font = bold_font
                    cell.border = border_total
                    cell.alignment = c_left if col_idx == 2 else c_center

                # Auto column width & freeze panes for Sheet 2
                ws2.freeze_panes = "A5"
                for col in range(1, 9):
                    max_len = len(sum_headers[col - 1]) + 2
                    for r_i in range(5, curr_row + 1):
                        cv = ws2.cell(row=r_i, column=col).value or ""
                        max_len = max(max_len, int(len(str(cv)) * 1.2))
                    col_letter = get_column_letter(col)
                    ws2.column_dimensions[col_letter].width = max(min(max_len, 35), 11)

                self._apply_print_setup(ws2, header_row=4)

                wb.save(filepath)

            messagebox.showinfo("Export Successful", f"Report saved successfully to:\n{filepath}")
            self.log(f"📥 Report exported to {filepath}", "SUCCESS")

            # Auto-open file if desktop environment allows
            if auto_open:
                try:
                    if sys.platform == "win32":
                        os.startfile(filepath)
                    elif sys.platform == "darwin":
                        import subprocess
                        subprocess.call(['open', filepath])
                except Exception:
                    pass

            return filepath
        except Exception as ex:
            messagebox.showerror("Export Failed", f"Could not save file: {ex}")
            self.log(f"❌ Export failed: {ex}", "ERROR")
            return None
